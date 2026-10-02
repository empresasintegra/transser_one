import re
from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_POST

from accounts.decorators import rol_requerido
from catalogos.models import Conductor, Tracto
from gastos.models import CategoriaGasto, GastoRuta
from maestros.models import Cliente, Local, Proveedor, Rampla, Ruta, Tarifa

from .models import ComisionServicio, EventoServicio, ParadaServicio, Servicio, TIPOS_TARIFA

ROL_CREADOR_SERVICIO = "Operaciones"
ROLES_AUTORIZADORES_CAMBIO = ("Administrador", "Gerencia")
ROL_MODIFICADOR_TARIFA = "Administrador"

ETIQUETAS_TIPO_TARIFA = dict(TIPOS_TARIFA)


def _resolver_tarifa(cliente_id, ruta_id, local_id=None):
    """Encuentra la Tarifa activa para (cliente, ruta), prefiriendo una
    específica del local elegido por sobre una genérica del cliente."""
    if not cliente_id or not ruta_id:
        return None
    base = Tarifa.objects.filter(cliente_id=cliente_id, ruta_id=ruta_id, activa=True)
    if local_id:
        con_local = base.filter(local_id=local_id).first()
        if con_local:
            return con_local
    return base.filter(local__isnull=True).first() or base.first()


@login_required
def dashboard(request):
    servicios = Servicio.objects.all()
    hoy = date.today()

    servicios_mes = servicios.filter(fecha_carga__year=hoy.year, fecha_carga__month=hoy.month)
    ventas_hoy = sum((s.tarifa for s in servicios.filter(fecha_carga=hoy)), Decimal("0"))
    ventas_mes = sum((s.tarifa for s in servicios_mes), Decimal("0"))
    facturacion_mes = sum((s.total_factura for s in servicios_mes), Decimal("0"))
    activos = servicios.exclude(estado=Servicio.Estado.FINALIZADO)
    servicios_mes_count = servicios_mes.count()
    ticket_promedio = (ventas_mes / servicios_mes_count) if servicios_mes_count else Decimal("0")

    contexto = {
        "ventas_hoy": ventas_hoy,
        "ventas_mes": ventas_mes,
        "facturacion_mes": facturacion_mes,
        "servicios_mes_count": servicios_mes_count,
        "activos_count": activos.count(),
        "finalizados_mes": servicios_mes.filter(estado=Servicio.Estado.FINALIZADO).count(),
        "clientes_mes": servicios_mes.values("cliente").distinct().count(),
        "ticket_promedio": ticket_promedio,
        "en_ruta": servicios.filter(estado=Servicio.Estado.EN_RUTA).count(),
        "distribucion": {etiqueta: servicios.filter(estado=valor).count() for valor, etiqueta in Servicio.Estado.choices},
        "recientes": servicios.select_related("cliente", "ruta__comuna_origen", "ruta__comuna_destino").order_by("-fecha_carga", "-id")[:8],
    }
    return render(request, "operaciones/dashboard.html", contexto)


@login_required
def centro_operaciones(request):
    servicios = Servicio.objects.select_related(
        "cliente", "ruta__comuna_origen", "ruta__comuna_destino",
        "conductor_principal", "tracto", "rampla",
    )
    busqueda = request.GET.get("q", "").strip()
    filtro_estado = request.GET.get("estado", "")

    if filtro_estado:
        servicios = servicios.filter(estado=filtro_estado)
    if busqueda:
        servicios = servicios.filter(
            Q(numero__icontains=busqueda)
            | Q(cliente__razon_social__icontains=busqueda)
            | Q(ruta__comuna_origen__nombre__icontains=busqueda)
            | Q(ruta__comuna_destino__nombre__icontains=busqueda)
            | Q(conductor_principal__nombre__icontains=busqueda)
            | Q(tracto__patente__icontains=busqueda)
            | Q(rampla__patente__icontains=busqueda)
        )

    resumen = [
        (valor, etiqueta, Servicio.objects.filter(estado=valor).count())
        for valor, etiqueta in Servicio.Estado.choices
    ]

    contexto = {
        "servicios": servicios,
        "estados": Servicio.Estado.choices,
        "resumen": resumen,
        "busqueda": busqueda,
        "filtro_estado": filtro_estado,
    }
    return render(request, "operaciones/centro_operaciones.html", contexto)


@login_required
def servicio_detalle(request, servicio_id):
    servicio = get_object_or_404(Servicio, pk=servicio_id)
    gastos = servicio.gastos.select_related("categoria", "proveedor").all()
    gastos_validos = gastos.exclude(estado=GastoRuta.Estado.RECHAZADO)
    gastos_netos = sum((g.monto_neto for g in gastos_validos), Decimal("0"))
    utilidad = servicio.tarifa - gastos_netos
    margen = (utilidad / servicio.tarifa * 100) if servicio.tarifa else Decimal("0")
    siguiente_estado = next(iter(Servicio.TRANSICIONES.get(servicio.estado, [])), None)

    cliente_servicio = servicio.cliente
    locales_cliente = Local.objects.filter(cliente=cliente_servicio).order_by("nombre")
    rutas_cliente = (
        Ruta.objects.filter(tarifas__cliente=cliente_servicio, tarifas__activa=True)
        .distinct().select_related("comuna_origen", "comuna_destino")
        .order_by("comuna_origen__nombre", "comuna_destino__nombre")
    )
    paradas = servicio.paradas.select_related("local", "ruta__comuna_origen", "ruta__comuna_destino")
    total_paradas = sum((p.valor for p in paradas), Decimal("0"))

    contexto = {
        "servicio": servicio,
        "gastos": gastos,
        "categorias": CategoriaGasto.objects.filter(activa=True),
        "gastos_netos": gastos_netos,
        "utilidad": utilidad,
        "margen": margen,
        "gastos_pendientes": gastos.filter(estado__in=[GastoRuta.Estado.PENDIENTE, GastoRuta.Estado.REVISADO]).count(),
        "siguiente_estado": siguiente_estado,
        "paradas": paradas,
        "total_paradas": total_paradas,
        "cliente_servicio": cliente_servicio,
        "locales_cliente": locales_cliente,
        "rutas_cliente": rutas_cliente,
        "ramplas": Rampla.objects.filter(activa=True),
        "conductores": Conductor.objects.filter(estado=Conductor.Estado.ACTIVO),
        "proveedores": Proveedor.objects.filter(activo=True),
    }
    return render(request, "operaciones/detalle.html", contexto)


@login_required
@require_POST
def crear_parada(request, servicio_id):
    servicio = get_object_or_404(Servicio, pk=servicio_id)
    cliente_servicio = servicio.cliente
    local_id = request.POST.get("local") or None
    if not local_id:
        messages.error(request, "Selecciona el local de la parada.")
        return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)
    local_obj = get_object_or_404(Local, pk=local_id, cliente=cliente_servicio)

    es_valor_extra = request.POST.get("es_valor_extra") == "on"
    observacion = request.POST.get("observacion") or None

    if es_valor_extra:
        try:
            valor = Decimal(request.POST.get("valor_extra") or "0")
        except Exception:
            valor = Decimal("0")
        if valor <= 0:
            messages.error(request, "Ingresa el valor a pagar en ese local.")
            return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)

        ParadaServicio.objects.create(
            servicio=servicio, local=local_obj, es_valor_extra=True,
            valor=valor, observacion=observacion, creado_por=request.user,
        )
        messages.success(request, "Parada con valor extra agregada.")
        return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)

    ruta_id = request.POST.get("ruta") or None
    tipo_tarifa = request.POST.get("tipo_tarifa") or ""
    if not ruta_id or tipo_tarifa not in ETIQUETAS_TIPO_TARIFA:
        messages.error(request, "Selecciona la ruta y el tipo de carga de la parada.")
        return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)

    tarifa_obj = _resolver_tarifa(cliente_servicio.id, ruta_id, local_id)
    valor = getattr(tarifa_obj, tipo_tarifa, None) if tarifa_obj else None
    if valor is None:
        messages.error(request, "No hay una tarifa activa con ese tipo de carga para ese local y esa ruta.")
        return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)

    ParadaServicio.objects.create(
        servicio=servicio, local=local_obj, es_valor_extra=False,
        ruta_id=ruta_id, tipo_tarifa=tipo_tarifa, tarifa_ref=tarifa_obj,
        valor=valor, observacion=observacion, creado_por=request.user,
    )
    messages.success(request, "Parada agregada.")
    return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)


@login_required
@require_POST
def eliminar_parada(request, parada_id):
    parada = get_object_or_404(ParadaServicio, pk=parada_id)
    servicio_id = parada.servicio_id
    parada.delete()
    messages.success(request, "Parada eliminada.")
    return redirect("operaciones:servicio_detalle", servicio_id=servicio_id)


@login_required
def opciones_rutas_cliente(request):
    """Fragmento htmx: rutas con al menos una Tarifa activa para el cliente elegido."""
    cliente_id = request.GET.get("cliente")
    rutas = Ruta.objects.none()
    if cliente_id:
        rutas = (
            Ruta.objects.filter(tarifas__cliente_id=cliente_id, tarifas__activa=True)
            .distinct()
            .select_related("comuna_origen", "comuna_destino")
            .order_by("comuna_origen__nombre", "comuna_destino__nombre")
        )
    return render(request, "operaciones/_opciones_ruta.html", {"rutas": rutas})


@login_required
def opciones_locales_cliente_servicio(request):
    """Fragmento htmx: locales del cliente elegido (campo opcional, a diferencia
    del mismo fragmento en el mantenedor de Tarifas)."""
    cliente_id = request.GET.get("cliente")
    locales = Local.objects.filter(cliente_id=cliente_id).order_by("nombre") if cliente_id else Local.objects.none()
    return render(request, "operaciones/_opciones_local.html", {"locales": locales})


@login_required
def previsualizar_tarifa_servicio(request):
    """Fragmento htmx: dado cliente + ruta (+ local opcional), resuelve la
    Tarifa activa y ofrece un <select> con los tipos de carga que tienen
    valor cargado (seco/frío/congelado/única) para elegir cuál corresponde."""
    tarifa_obj = _resolver_tarifa(
        request.GET.get("cliente"), request.GET.get("ruta"), request.GET.get("local")
    )
    opciones = []
    if tarifa_obj:
        for campo, etiqueta in TIPOS_TARIFA:
            valor = getattr(tarifa_obj, campo)
            if valor is not None:
                opciones.append((campo, etiqueta, valor))
    contexto = {
        "ruta_elegida": bool(request.GET.get("ruta")),
        "tarifa": tarifa_obj,
        "opciones": opciones,
    }
    return render(request, "operaciones/_tarifa_preview.html", contexto)


@rol_requerido(ROL_CREADOR_SERVICIO)
def crear_servicio(request):
    if request.method == "POST":
        cliente_id = request.POST.get("cliente") or None
        local_id = request.POST.get("local") or None
        ruta_id = request.POST.get("ruta") or None
        tipo_tarifa = request.POST.get("tipo_tarifa") or ""
        fecha_carga = parse_date(request.POST.get("fecha_carga") or "")

        if not cliente_id or not ruta_id or not fecha_carga:
            messages.error(request, "Completa el cliente, la ruta y la fecha de carga.")
            return redirect("operaciones:crear_servicio")

        if tipo_tarifa not in ETIQUETAS_TIPO_TARIFA:
            messages.error(request, "Selecciona el tipo de tarifa (seco, frío, congelado o única).")
            return redirect("operaciones:crear_servicio")

        cliente_obj = get_object_or_404(Cliente, pk=cliente_id)
        local_obj = get_object_or_404(Local, pk=local_id, cliente=cliente_obj) if local_id else None
        tarifa_obj = _resolver_tarifa(cliente_id, ruta_id, local_id)
        valor_tarifa = getattr(tarifa_obj, tipo_tarifa, None) if tarifa_obj else None

        if valor_tarifa is None:
            messages.error(request, "No hay una tarifa activa con ese tipo de carga para el cliente y la ruta elegidos.")
            return redirect("operaciones:crear_servicio")

        conductor_id = request.POST.get("conductor_principal") or None
        conductor_obj = get_object_or_404(Conductor, pk=conductor_id, estado=Conductor.Estado.ACTIVO) if conductor_id else None
        rampla_id = request.POST.get("rampla") or None
        rampla_obj = get_object_or_404(Rampla, pk=rampla_id, activa=True) if rampla_id else None

        servicio = Servicio(
            cliente=cliente_obj,
            local_destino=local_obj,
            tarifa_ref=tarifa_obj,
            ruta=tarifa_obj.ruta,
            fecha_carga=fecha_carga,
            fecha_entrega_estimada=parse_date(request.POST.get("fecha_entrega_estimada") or ""),
            tipo_carga=tipo_tarifa,
            conductor_principal=conductor_obj,
            tracto=Tracto.objects.filter(conductor=conductor_obj).first() if conductor_obj else None,
            rampla=rampla_obj,
            tarifa=valor_tarifa,
            tasa_iva=Decimal("19"),
            observaciones=request.POST.get("observaciones") or None,
            creado_por=request.user,
        )
        servicio.calcular_iva_total()
        servicio.save()
        servicio.numero = f"TS-{servicio.fecha_carga.year}-{servicio.id:06d}"
        servicio.save(update_fields=["numero"])

        EventoServicio.objects.create(
            servicio=servicio,
            tipo="CREACION",
            estado_nuevo=Servicio.Estado.PROGRAMADO,
            actor=request.user,
            motivo="Creación inicial del servicio.",
        )
        if servicio.conductor_principal:
            ComisionServicio.objects.create(
                servicio=servicio,
                conductor=servicio.conductor_principal,
                porcentaje=Decimal("100.00"),
                motivo="Conductor principal del servicio.",
            )

        messages.success(request, f"Servicio {servicio.numero} creado correctamente.")
        return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)

    contexto = {
        "clientes": Cliente.objects.filter(activo=True).order_by("razon_social"),
        "conductores": Conductor.objects.filter(estado=Conductor.Estado.ACTIVO),
        "ramplas": Rampla.objects.filter(activa=True),
    }
    return render(request, "operaciones/servicio_form.html", contexto)


@login_required
@require_POST
def cambiar_estado(request, servicio_id):
    servicio = get_object_or_404(Servicio, pk=servicio_id)
    nuevo_estado = request.POST.get("nuevo_estado", "").strip()
    evidencia_foto = request.POST.get("evidencia_foto") or None
    motivo = request.POST.get("motivo") or f"Avance operacional a {nuevo_estado}"

    permitidos = Servicio.TRANSICIONES.get(servicio.estado, set())
    if nuevo_estado not in permitidos:
        messages.error(request, f"No se puede pasar de '{servicio.estado}' a '{nuevo_estado}'.")
        return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)

    if nuevo_estado in {Servicio.Estado.EN_RUTA, Servicio.Estado.FINALIZADO} and not evidencia_foto:
        messages.error(request, f"Para pasar a '{nuevo_estado}' es obligatoria una fotografía.")
        return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)

    estado_anterior = servicio.estado
    servicio.estado = nuevo_estado
    servicio.save(update_fields=["estado", "actualizado_en"])
    EventoServicio.objects.create(
        servicio=servicio,
        tipo="CAMBIO_ESTADO",
        estado_anterior=estado_anterior,
        estado_nuevo=nuevo_estado,
        actor=request.user,
        motivo=motivo,
        evidencia_foto=evidencia_foto,
    )
    messages.success(request, f"Servicio avanzado a {nuevo_estado}.")
    return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)


@login_required
@require_POST
def cambiar_rampla(request, servicio_id):
    servicio = get_object_or_404(Servicio, pk=servicio_id)
    if servicio.estado not in {Servicio.Estado.EN_RUTA, Servicio.Estado.DESCARGANDO}:
        messages.error(request, "El cambio de rampla solo se permite En Ruta o Descargando.")
        return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)

    rampla_nueva = Rampla.objects.filter(pk=request.POST.get("rampla_nueva") or None, activa=True).first()
    conductor_secundario = Conductor.objects.filter(
        pk=request.POST.get("conductor_secundario") or None, estado=Conductor.Estado.ACTIVO,
    ).first()
    motivo = request.POST.get("motivo", "").strip()
    if not rampla_nueva or not conductor_secundario or len(motivo) < 3:
        messages.error(request, "Rampla, conductor secundario y motivo son obligatorios.")
        return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)

    rampla_anterior = servicio.rampla
    servicio.rampla = rampla_nueva
    servicio.conductor_secundario = conductor_secundario
    servicio.save(update_fields=["rampla", "conductor_secundario", "actualizado_en"])

    servicio.comisiones.all().delete()
    if servicio.conductor_principal:
        ComisionServicio.objects.create(
            servicio=servicio, conductor=servicio.conductor_principal,
            porcentaje=Decimal("50.00"), motivo="Reparto por reemplazo de rampla durante el viaje.",
        )
    ComisionServicio.objects.create(
        servicio=servicio, conductor=conductor_secundario,
        porcentaje=Decimal("50.00"), motivo="Reparto por reemplazo de rampla durante el viaje.",
    )

    EventoServicio.objects.create(
        servicio=servicio, tipo="CAMBIO_RAMPLA", rampla_anterior=rampla_anterior,
        rampla_nueva=rampla_nueva, actor=request.user, motivo=motivo,
    )
    messages.success(request, "Rampla actualizada.")
    return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)


@rol_requerido(ROL_MODIFICADOR_TARIFA)
@require_POST
def modificar_tarifa(request, servicio_id):
    servicio = get_object_or_404(Servicio, pk=servicio_id)
    motivo = request.POST.get("motivo", "").strip()
    try:
        tarifa_nueva = Decimal(request.POST.get("tarifa_nueva", "0"))
    except Exception:
        tarifa_nueva = Decimal("0")

    if tarifa_nueva <= 0 or len(motivo) < 3:
        messages.error(request, "Tarifa y motivo son obligatorios.")
        return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)

    tarifa_anterior = servicio.tarifa
    servicio.tarifa = tarifa_nueva
    servicio.calcular_iva_total()
    servicio.save(update_fields=["tarifa", "iva", "total_factura", "actualizado_en"])
    EventoServicio.objects.create(
        servicio=servicio, tipo="CAMBIO_TARIFA", actor=request.user,
        motivo=f"{motivo} Tarifa anterior: {tarifa_anterior}. Tarifa nueva: {tarifa_nueva}.",
    )
    messages.success(request, "Tarifa actualizada.")
    return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)


@rol_requerido(*ROLES_AUTORIZADORES_CAMBIO)
@require_POST
def modificar_datos(request, servicio_id):
    servicio = get_object_or_404(Servicio, pk=servicio_id)
    motivo = request.POST.get("motivo", "").strip()
    if len(motivo) < 3:
        messages.error(request, "El motivo es obligatorio.")
        return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)

    ruta_id = request.POST.get("ruta")
    if ruta_id:
        servicio.ruta = get_object_or_404(Ruta, pk=ruta_id)
    for campo in ("fecha_carga", "fecha_entrega_estimada"):
        valor = parse_date(request.POST.get(campo) or "")
        if valor:
            setattr(servicio, campo, valor)
    servicio.save()

    EventoServicio.objects.create(
        servicio=servicio, tipo="CAMBIO_DATOS", actor=request.user,
        autorizado_por=request.user, motivo=motivo,
    )
    messages.success(request, "Datos del servicio actualizados.")
    return redirect("operaciones:servicio_detalle", servicio_id=servicio.id)


# --- Reportes ------------------------------------------------------------

ENCABEZADOS_SERVICIOS = [
    "N° Servicio", "Fecha de carga", "Entrega estimada", "Cliente", "Origen", "Destino",
    "Conductor principal", "Conductor secundario", "Tracto", "Rampla", "Tipo de carga",
    "Tarifa neta", "IVA", "Total factura", "Estado", "Creado por", "Observaciones",
    # Resumen de paradas directo en esta hoja (el detalle fila por fila
    # queda en la hoja "Paradas") — para que no dependa de que alguien
    # note que hay una segunda pestaña al abrir el archivo.
    "N° Paradas", "Valor paradas normales", "Valor paradas extra", "Total paradas",
]

ENCABEZADOS_PARADAS = [
    "N° Servicio", "Local", "Modalidad", "Ruta", "Tipo de carga", "Valor", "Observación",
]


def _hoja_con_encabezado(wb, titulo, encabezados):
    ws = wb.create_sheet(titulo)
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    return ws


def _construir_excel_servicios(cliente, servicios):
    """Arma el .xlsx: una hoja "Servicios" (una fila por servicio) y una
    hoja "Paradas" (una fila por parada, con el N° de servicio para
    cruzarlas) — una parada por fila es más claro en una planilla que
    amontonar varias en una sola celda."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_servicios = _hoja_con_encabezado(wb, "Servicios", ENCABEZADOS_SERVICIOS)
    ws_paradas = _hoja_con_encabezado(wb, "Paradas", ENCABEZADOS_PARADAS)

    for s in servicios:
        paradas_servicio = list(s.paradas.all())
        valor_normales = sum((p.valor for p in paradas_servicio if not p.es_valor_extra), Decimal("0"))
        valor_extra = sum((p.valor for p in paradas_servicio if p.es_valor_extra), Decimal("0"))

        ws_servicios.append([
            s.numero, s.fecha_carga, s.fecha_entrega_estimada, s.cliente.razon_social, s.origen, s.destino,
            s.conductor_principal.nombre if s.conductor_principal_id else None,
            s.conductor_secundario.nombre if s.conductor_secundario_id else None,
            s.tracto.patente if s.tracto_id else None,
            s.rampla.patente if s.rampla_id else None,
            s.get_tipo_carga_display() if s.tipo_carga else None,
            s.tarifa, s.iva, s.total_factura, s.estado, s.creado_por.nombre, s.observaciones,
            len(paradas_servicio), valor_normales, valor_extra, valor_normales + valor_extra,
        ])
        for p in paradas_servicio:
            ws_paradas.append([
                s.numero,
                p.local.nombre if p.local_id else None,
                "Valor extra" if p.es_valor_extra else "Normal",
                str(p.ruta) if p.ruta_id else None,
                p.get_tipo_tarifa_display() if p.tipo_tarifa else None,
                p.valor,
                p.observacion,
            ])

    for ws in (ws_servicios, ws_paradas):
        for columna in ws.columns:
            largo = max((len(str(c.value)) for c in columna if c.value is not None), default=10)
            ws.column_dimensions[columna[0].column_letter].width = min(largo + 2, 45)

    return wb


@login_required
def reportes(request):
    contexto = {"clientes": Cliente.objects.filter(activo=True).order_by("razon_social")}
    return render(request, "operaciones/reportes.html", contexto)


@login_required
def reportes_excel(request):
    cliente_id = request.GET.get("cliente")
    mes = request.GET.get("mes", "")

    cliente = get_object_or_404(Cliente, pk=cliente_id) if cliente_id else None
    coincide_mes = re.fullmatch(r"(\d{4})-(\d{2})", mes)
    if not cliente or not coincide_mes:
        messages.error(request, "Selecciona un cliente y un mes para descargar el reporte.")
        return redirect("operaciones:reportes")

    anio, mes_numero = int(coincide_mes.group(1)), int(coincide_mes.group(2))
    servicios = (
        Servicio.objects.filter(cliente=cliente, fecha_carga__year=anio, fecha_carga__month=mes_numero)
        .select_related(
            "cliente", "ruta__comuna_origen", "ruta__comuna_destino", "conductor_principal",
            "conductor_secundario", "tracto", "rampla", "creado_por",
        )
        .prefetch_related("paradas__local", "paradas__ruta__comuna_origen", "paradas__ruta__comuna_destino")
        .order_by("fecha_carga", "numero")
    )

    wb = _construir_excel_servicios(cliente, servicios)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_cliente = re.sub(r"[^A-Za-z0-9]+", "-", cliente.razon_social).strip("-")
    nombre_archivo = f"Servicios_{nombre_cliente}_{mes}.xlsx"
    respuesta = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre_archivo}"'
    return respuesta


# --- Guía de despacho (PDF) ----------------------------------------------
# Replica el formato de guía de despacho que la empresa ya usa en papel,
# con los datos que SÍ tenemos en el sistema (cliente, local, ruta,
# conductor, tracto, rampla, tarifa, paradas). A propósito NO incluye
# timbre electrónico SII ni numeración de folio/CAF: no estamos
# integrados con el SII, así que sería un documento tributario falso.
# Esto es un comprobante operativo interno — reemplaza el papel que hoy
# llena el conductor a mano, no una Guía de Despacho Electrónica SII.

_COLOR_MARINO = colors.HexColor("#101c2c")
_COLOR_NARANJA = colors.HexColor("#f26716")
_COLOR_TEXTO = colors.HexColor("#1f2933")
_COLOR_GRIS = colors.HexColor("#6b7686")
_COLOR_BORDE = colors.HexColor("#d9dee5")
_COLOR_FONDO = colors.HexColor("#f4f6f8")

# Todo el texto de las tablas va en Paragraph (no strings sueltos): así
# reportlab lo ajusta al ancho de la celda en vez de dejarlo salirse.
_ESTILOS = getSampleStyleSheet()
_BASE = ParagraphStyle("guia_base", parent=_ESTILOS["Normal"], fontName="Helvetica", fontSize=9, leading=12, textColor=_COLOR_TEXTO)
_E_ETIQUETA = ParagraphStyle("guia_etiqueta", parent=_BASE, fontSize=7, leading=9, textColor=_COLOR_GRIS)
_E_VALOR = ParagraphStyle("guia_valor", parent=_BASE, fontName="Helvetica-Bold", fontSize=9, leading=11)
_E_TITULO_SECCION = ParagraphStyle("guia_seccion", parent=_BASE, fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.white)
_E_MARCA = ParagraphStyle("guia_marca", parent=_BASE, fontName="Helvetica-Bold", fontSize=20, leading=22, textColor=colors.white)
_E_MARCA_SUB = ParagraphStyle("guia_marca_sub", parent=_BASE, fontSize=8, leading=10, textColor=colors.HexColor("#b8c2cf"))
_E_FOLIO_TIT = ParagraphStyle("guia_folio_tit", parent=_BASE, fontName="Helvetica-Bold", fontSize=9, leading=11, alignment=1, textColor=_COLOR_NARANJA)
_E_FOLIO_NUM = ParagraphStyle("guia_folio_num", parent=_BASE, fontName="Helvetica-Bold", fontSize=13, leading=16, alignment=1)
_E_FOLIO_EST = ParagraphStyle("guia_folio_est", parent=_BASE, fontSize=8, leading=10, alignment=1, textColor=_COLOR_GRIS)
_E_TH = ParagraphStyle("guia_th", parent=_BASE, fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.white)
_E_TH_DER = ParagraphStyle("guia_th_der", parent=_E_TH, alignment=2)
_E_CELDA = ParagraphStyle("guia_celda", parent=_BASE, fontSize=9, leading=11)
_E_CELDA_DER = ParagraphStyle("guia_celda_der", parent=_E_CELDA, alignment=2)
_E_TOT_ETIQ = ParagraphStyle("guia_tot_etiq", parent=_BASE, fontSize=9, leading=11, alignment=2, textColor=_COLOR_GRIS)
_E_TOT_VAL = ParagraphStyle("guia_tot_val", parent=_BASE, fontName="Helvetica-Bold", fontSize=9, leading=11, alignment=2)
_E_TOT_FINAL_ETIQ = ParagraphStyle("guia_tot_f_etiq", parent=_BASE, fontName="Helvetica-Bold", fontSize=10, leading=12, alignment=2, textColor=colors.white)
_E_TOT_FINAL_VAL = ParagraphStyle("guia_tot_f_val", parent=_E_TOT_FINAL_ETIQ, fontSize=12, leading=14)
_E_LEGAL = ParagraphStyle("guia_legal", parent=_BASE, fontSize=7, leading=9, textColor=_COLOR_GRIS)
_E_FIRMA = ParagraphStyle("guia_firma", parent=_BASE, fontSize=8, leading=10, textColor=_COLOR_GRIS)

_ANCHO = letter[0] - 3 * cm  # ancho útil con márgenes de 1,5 cm


def _clp(valor):
    return "$" + f"{valor:,.0f}".replace(",", ".")


def _esc(texto):
    """Escapa &, < y > para que un valor con esos caracteres no rompa el markup de Paragraph."""
    return str(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _campo(etiqueta, valor):
    """Etiqueta chica gris arriba + valor en negrita abajo, en una sola celda."""
    return [Paragraph(_esc(etiqueta).upper(), _E_ETIQUETA), Paragraph(_esc(valor) if valor else "—", _E_VALOR)]


def _tarjeta(titulo, campos, ancho):
    """Bloque con barra de título azul marino y los campos en filas."""
    filas = [[Paragraph(_esc(titulo).upper(), _E_TITULO_SECCION)]] + [[c] for c in campos]
    tabla = Table(filas, colWidths=[ancho])
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), _COLOR_MARINO),
        ("BOX", (0, 0), (-1, -1), 0.6, _COLOR_BORDE),
        ("LINEBELOW", (0, 1), (-1, -2), 0.4, _COLOR_BORDE),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 5),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
        ("TOPPADDING", (0, 1), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 3),
    ]))
    return tabla


def _pie_de_pagina(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(_COLOR_BORDE)
    canvas.line(1.5 * cm, 1.4 * cm, letter[0] - 1.5 * cm, 1.4 * cm)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(_COLOR_GRIS)
    canvas.drawString(1.5 * cm, 1.0 * cm, "Documento generado por Transser One - comprobante operativo interno. No constituye Guía de Despacho Electrónica ante el SII.")
    canvas.drawRightString(letter[0] - 1.5 * cm, 1.0 * cm, f"Página {doc.page}")
    canvas.restoreState()


def _construir_pdf_guia(servicio):
    cliente = servicio.cliente
    conductor = servicio.conductor_principal
    local = servicio.local_destino
    paradas = list(servicio.paradas.select_related("local", "ruta__comuna_origen", "ruta__comuna_destino"))

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter,
        topMargin=1.5 * cm, bottomMargin=2.0 * cm, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        title=f"Guía de despacho {servicio.numero}", author="Transser One",
    )
    el = []

    # --- Encabezado: marca a la izquierda, folio a la derecha ---
    ancho_folio = 6.2 * cm
    folio = Table(
        [[Paragraph("GUÍA DE DESPACHO", _E_FOLIO_TIT)],
         [Paragraph(f"N° {_esc(servicio.numero)}", _E_FOLIO_NUM)],
         [Paragraph(f"Emisión: {servicio.fecha_carga.strftime('%d-%m-%Y')} · Estado: {_esc(servicio.estado)}", _E_FOLIO_EST)]],
        colWidths=[ancho_folio],
    )
    folio.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("BOX", (0, 0), (-1, -1), 1.4, _COLOR_NARANJA),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    cabecera = Table(
        [[[Paragraph("TRANSSER ONE", _E_MARCA), Paragraph("Transporte y logística", _E_MARCA_SUB)], folio]],
        colWidths=[_ANCHO - ancho_folio - 0.6 * cm, ancho_folio + 0.6 * cm],
    )
    cabecera.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), _COLOR_MARINO),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (0, 0), 14),
        ("RIGHTPADDING", (1, 0), (1, 0), 14),
        ("TOPPADDING", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ("LINEBELOW", (0, 0), (-1, -1), 3, _COLOR_NARANJA),
    ]))
    el += [cabecera, Spacer(1, 0.45 * cm)]

    # --- Origen (cliente) y destino lado a lado ---
    mitad = (_ANCHO - 0.5 * cm) / 2
    origen = _tarjeta("Datos origen (cliente)", [
        _campo("Razón social", cliente.razon_social),
        _campo("RUT", cliente.rut),
        _campo("Giro", cliente.giro),
        _campo("Dirección", cliente.direccion),
        _campo("Comuna de origen", servicio.origen),
    ], mitad)
    destino = _tarjeta("Datos destino", [
        _campo("Local", local.nombre if local else None),
        _campo("Código de local", servicio.codigo_local),
        _campo("Dirección", local.direccion if local else None),
        _campo("Comuna de destino", servicio.destino),
        _campo("Entrega estimada", servicio.fecha_entrega_estimada.strftime("%d-%m-%Y") if servicio.fecha_entrega_estimada else None),
    ], mitad)
    pareja = Table([[origen, "", destino]], colWidths=[mitad, 0.5 * cm, mitad])
    pareja.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    el += [pareja, Spacer(1, 0.4 * cm)]

    # --- Transporte: cuadrícula de 4 columnas ---
    cuarto = _ANCHO / 4
    transporte = Table(
        [[Paragraph("DATOS DE TRANSPORTE", _E_TITULO_SECCION), "", "", ""],
         [_campo("Transportista", conductor.proveedor.nombre if conductor and conductor.proveedor_id else None),
          _campo("Conductor", conductor.nombre if conductor else None),
          _campo("RUT conductor", conductor.rut if conductor else None),
          _campo("Tipo de carga", servicio.get_tipo_carga_display() if servicio.tipo_carga else None)],
         [_campo("Patente tracto", servicio.tracto.patente if servicio.tracto_id else None),
          _campo("Patente rampla", servicio.rampla.patente if servicio.rampla_id else None),
          _campo("Segundo conductor", servicio.conductor_secundario.nombre if servicio.conductor_secundario_id else None),
          _campo("Paradas adicionales", str(len(paradas)))]],
        colWidths=[cuarto] * 4,
    )
    transporte.setStyle(TableStyle([
        ("SPAN", (0, 0), (-1, 0)),
        ("BACKGROUND", (0, 0), (-1, 0), _COLOR_MARINO),
        ("BOX", (0, 0), (-1, -1), 0.6, _COLOR_BORDE),
        ("LINEBELOW", (0, 1), (-1, 1), 0.4, _COLOR_BORDE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 5),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
        ("TOPPADDING", (0, 1), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
    ]))
    el += [transporte, Spacer(1, 0.4 * cm)]

    # --- Ítems: viaje principal + paradas (texto en Paragraph: se ajusta al ancho) ---
    filas = [[Paragraph("ÍTEM", _E_TH), Paragraph("DESCRIPCIÓN", _E_TH), Paragraph("MONTO", _E_TH_DER)]]
    filas.append([
        Paragraph("Transporte", _E_CELDA),
        Paragraph(f"{_esc(servicio.origen)} → {_esc(servicio.destino)}<br/>"
                  f"<font size=7 color='#6b7686'>Carga: {_esc(servicio.get_tipo_carga_display() or 'Sin tipo')}</font>", _E_CELDA),
        Paragraph(_clp(servicio.tarifa), _E_CELDA_DER),
    ])
    for p in paradas:
        if p.es_valor_extra:
            detalle = "Valor extra"
        else:
            detalle = f"{_esc(p.ruta)}<br/><font size=7 color='#6b7686'>Carga: {_esc(p.get_tipo_tarifa_display() or '—')}</font>"
        filas.append([
            Paragraph(f"Parada: {_esc(p.local.nombre)}", _E_CELDA),
            Paragraph(detalle + (f"<br/><font size=7 color='#6b7686'>{_esc(p.observacion)}</font>" if p.observacion else ""), _E_CELDA),
            Paragraph(_clp(p.valor), _E_CELDA_DER),
        ])
    ancho_monto = 3.4 * cm
    ancho_item = 5.2 * cm
    items = Table(filas, colWidths=[ancho_item, _ANCHO - ancho_item - ancho_monto, ancho_monto], repeatRows=1)
    estilo_items = [
        ("BACKGROUND", (0, 0), (-1, 0), _COLOR_MARINO),
        ("BOX", (0, 0), (-1, -1), 0.6, _COLOR_BORDE),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, _COLOR_BORDE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]
    for i in range(2, len(filas), 2):
        estilo_items.append(("BACKGROUND", (0, i), (-1, i), _COLOR_FONDO))
    items.setStyle(TableStyle(estilo_items))
    el += [items, Spacer(1, 0.3 * cm)]

    # --- Totales alineados a la derecha ---
    # Los totales suman TODO lo listado arriba (viaje + paradas), no solo el
    # viaje principal: si no, el documento no cuadraría con sus propios ítems.
    neto = servicio.tarifa + sum((p.valor for p in paradas), Decimal("0"))
    iva = (neto * servicio.tasa_iva / Decimal("100")).quantize(Decimal("1"))
    ancho_tot = 8 * cm
    totales = Table(
        [[Paragraph("Neto", _E_TOT_ETIQ), Paragraph(_clp(neto), _E_TOT_VAL)],
         [Paragraph(f"IVA ({servicio.tasa_iva:.0f}%)", _E_TOT_ETIQ), Paragraph(_clp(iva), _E_TOT_VAL)],
         [Paragraph("TOTAL", _E_TOT_FINAL_ETIQ), Paragraph(_clp(neto + iva), _E_TOT_FINAL_VAL)]],
        colWidths=[ancho_tot * 0.45, ancho_tot * 0.55],
    )
    totales.setStyle(TableStyle([
        ("BACKGROUND", (0, 2), (-1, 2), _COLOR_NARANJA),
        ("LINEBELOW", (0, 0), (-1, 0), 0.4, _COLOR_BORDE),
        ("BOX", (0, 0), (-1, -1), 0.6, _COLOR_BORDE),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    contenedor = Table([["", totales]], colWidths=[_ANCHO - ancho_tot, ancho_tot])
    contenedor.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    el += [contenedor, Spacer(1, 0.5 * cm)]

    # --- Acuse de recibo + firma (KeepTogether: no se parte entre páginas) ---
    linea = lambda t: Paragraph(t, _E_FIRMA)  # noqa: E731
    firma = Table(
        [[linea("Nombre"), "", linea("RUT"), ""],
         [linea("Recinto"), "", linea("Fecha"), ""],
         [linea("Firma"), "", "", ""]],
        colWidths=[2.0 * cm, (_ANCHO - 6.0 * cm) / 2, 1.6 * cm, (_ANCHO - 6.0 * cm) / 2 + 2.4 * cm],
        rowHeights=[0.8 * cm, 0.8 * cm, 1.2 * cm],
    )
    firma.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ("LINEBELOW", (1, 0), (1, 0), 0.5, _COLOR_TEXTO),
        ("LINEBELOW", (3, 0), (3, 0), 0.5, _COLOR_TEXTO),
        ("LINEBELOW", (1, 1), (1, 1), 0.5, _COLOR_TEXTO),
        ("LINEBELOW", (3, 1), (3, 1), 0.5, _COLOR_TEXTO),
        ("LINEBELOW", (1, 2), (3, 2), 0.5, _COLOR_TEXTO),
    ]))
    el.append(KeepTogether([
        Paragraph("<b>ACUSE DE RECIBO</b>", _E_ETIQUETA),
        Spacer(1, 2),
        Paragraph(
            "El acuse de recibo que se declara en este acto acredita que la entrega de mercaderías o "
            "servicio(s) prestado(s) ha(n) sido recibido(s) conforme.", _E_LEGAL),
        Spacer(1, 0.2 * cm),
        firma,
    ]))

    doc.build(el, onFirstPage=_pie_de_pagina, onLaterPages=_pie_de_pagina)
    buffer.seek(0)
    return buffer


@login_required
def guia_despacho_pdf(request, servicio_id):
    servicio = get_object_or_404(
        Servicio.objects.select_related(
            "cliente", "local_destino", "ruta__comuna_origen", "ruta__comuna_destino",
            "conductor_principal__proveedor", "conductor_secundario", "tracto", "rampla",
        ),
        pk=servicio_id,
    )
    buffer = _construir_pdf_guia(servicio)
    respuesta = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    respuesta["Content-Disposition"] = f'inline; filename="Guia_despacho_{servicio.numero}.pdf"'
    return respuesta
